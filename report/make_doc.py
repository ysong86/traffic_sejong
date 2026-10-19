# -*- coding: utf-8 -*-
"""그림 제목·설명을 워드 파일로 뽑는다.

그림(PNG)에는 제목이 들어 있지 않다. 보고서 본문에서 캡션을 달기 때문에
그림이 제목을 품고 있으면 제목이 두 번 나온다. 그래서 캡션에 쓸 글을 여기
모아 둔다 — 파일 이름을 같이 적어 두어 그림과 짝짓기 쉽게 했다.

    C:\\Python313\\python.exe -I report/make_doc.py
      → report/세종시_교통안전_기본계획_그림목록.docx
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import base
import index_page

FONT = "맑은 고딕"
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK2 = RGBColor(0x52, 0x51, 0x4E)
INK3 = RGBColor(0x8A, 0x88, 0x80)
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)


def _hangul(run):
    """한글 글꼴은 동아시아 글꼴 속성을 따로 지정해야 먹는다."""
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def para(doc, text, size=10.5, bold=False, color=INK, space_after=4,
         align=None, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.35
    if indent is not None:
        pf.left_indent = Cm(indent)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _hangul(run)
    return p


def rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    pr = p._p.get_or_add_pPr()
    borders = pr.makeelement(qn("w:pBdr"), {})
    bottom = pr.makeelement(qn("w:bottom"), {})
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"),
                 ("w:color", "D9D8D3")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    pr.append(borders)


def build(made, stamp, counts):
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.4)
        section.top_margin = section.bottom_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    para(doc, "세종시 교통안전 기본계획 — 분석 그림 목록", 18, True, INK, 4)
    para(doc, "그림 %d장 · 자료 기준 %s · 사고다발지 %d곳 · 안전시설 %d개소"
         % (len(made), stamp, counts[0], counts[1]), 9.5, False, INK3, 2)
    para(doc, "작성 세종연구원 도시안전연구센터 · 문의 ayj826@sri.re.kr",
         9.5, False, INK3, 14)

    para(doc, "그림 파일은 report/figures/ 에 PNG(200dpi)로 있습니다. "
              "그림에는 제목이 들어 있지 않으므로 아래 제목·설명을 보고서 "
              "본문의 캡션으로 쓰시면 됩니다.", 10, False, INK2, 6)
    para(doc, "‘사업 근거’는 그 그림이 어떤 사업을 뒷받침하는지 적은 것입니다. "
              "기본계획에서 그림이 하는 일은 사업을 정당화하는 것이라서 "
              "같이 적어 두었습니다.", 10, False, INK2, 16)

    for m in made:
        use = index_page.USE.get(index_page.slug(m["file"]))
        rule(doc)
        para(doc, m["title"], 13, True, INK, 3)
        para(doc, "파일  %s" % m["file"], 9, False, INK3, 6)
        para(doc, m["subtitle"], 10.5, False, INK, 6)
        if use:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.left_indent = Cm(0.4)
            r1 = p.add_run("사업 근거  ")
            r1.font.size = Pt(10)
            r1.font.bold = True
            r1.font.color.rgb = ACCENT
            _hangul(r1)
            r2 = p.add_run("%s — %s" % (use[0], use[1]))
            r2.font.size = Pt(10)
            r2.font.color.rgb = INK2
            _hangul(r2)
        if m.get("source"):
            para(doc, m["source"], 8.5, False, INK3, 12)

    rule(doc)
    para(doc, "읽기 전에 아셔야 할 것", 13, True, INK, 6)
    for text in [
        "① 사고 자료는 두 가지가 섞여 있습니다. ‘지자체별 교통사고 통계’는 "
        "세종시 전체 사고(연도별 추이 그림)이고, ‘교통사고 다발지역’은 기준을 "
        "넘은 지점만 추린 것(지점 그림)입니다. 두 숫자를 더하거나 비교하면 "
        "안 됩니다.",
        "② 시설 자료에 지역 편중이 있습니다. 횡단보도·도로안전표지는 세종시가 "
        "등록한 신도심 분만 있어 조치원읍이 0건으로 나옵니다. 현장에 시설이 "
        "없다는 뜻이 아닙니다.",
        "③ 사망자 0명으로 보이는 지점 그림은 다발지역 지정 기준에 걸린 13곳 "
        "안에서 사망사고가 없었다는 뜻이며, 세종시 사망자가 0명이라는 뜻이 "
        "아닙니다(2025년 10명).",
    ]:
        para(doc, text, 10, False, INK2, 8, indent=0.4)

    path = os.path.join(base.HERE, "세종시_교통안전_기본계획_그림목록.docx")
    doc.save(path)
    return path


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    import make_figures
    import prep
    P = prep.build()
    base.MADE.clear()
    for fn in make_figures.FIGURES:
        fn(P)
    path = build(base.MADE, (P["raw"].get("generated_at") or "")[:10],
                 (len(P["spots"]), len(P["items"])))
    print("  %-40s %4d KB" % (os.path.basename(path),
                              os.path.getsize(path) // 1024))


if __name__ == "__main__":
    main()
